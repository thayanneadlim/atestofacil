import zipfile
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from lxml import etree
from PIL import Image

from servicos.models import (
    calcular_valor_dia,
    montar_resumo_pagamento,
)


def _png_minimo():
    imagem = Image.new("RGB", (20, 20), "white")
    buffer = BytesIO()
    imagem.save(buffer, format="PNG")
    return SimpleUploadedFile("atesto.png", buffer.getvalue(), content_type="image/png")


def _dia(horas, hora_fim, hora_inicio=time(8, 0), data=date(2026, 10, 1)):
    return SimpleNamespace(
        data=data,
        hora_inicio=hora_inicio,
        hora_fim=hora_fim,
        quantidade_horas=Decimal(str(horas)),
    )


class CalculoHorasTests(SimpleTestCase):
    def test_turno_que_passa_da_meia_noite(self):
        from django.core.exceptions import ValidationError

        from servicos.models import calcular_quantidade_horas

        self.assertEqual(
            calcular_quantidade_horas(time(18, 0), time(6, 0)),
            Decimal("12.00"),
        )
        self.assertEqual(
            calcular_quantidade_horas(time(22, 30), time(6, 0)),
            Decimal("7.50"),
        )
        with self.assertRaises(ValidationError):
            calcular_quantidade_horas(time(8, 0), time(8, 0))


class CalculoValorDiaTests(SimpleTestCase):
    def test_faixa_6_a_8_e_4_horas(self):
        self.assertEqual(calcular_valor_dia(8, time(17, 0)), Decimal("110.00"))
        self.assertEqual(calcular_valor_dia(7, time(16, 0)), Decimal("110.00"))
        self.assertEqual(calcular_valor_dia(6, time(14, 0)), Decimal("110.00"))
        self.assertEqual(calcular_valor_dia(10, time(18, 0)), Decimal("110.00"))
        self.assertEqual(calcular_valor_dia(4, time(12, 0)), Decimal("55.00"))
        self.assertEqual(calcular_valor_dia(5, time(13, 0)), Decimal("55.00"))

    def test_termino_as_22_nao_e_apos_22(self):
        self.assertEqual(calcular_valor_dia(8, time(22, 0)), Decimal("110.00"))
        self.assertEqual(calcular_valor_dia(4, time(22, 0)), Decimal("55.00"))
        self.assertEqual(calcular_valor_dia(5, time(22, 0)), Decimal("55.00"))

    def test_termino_depois_das_22(self):
        self.assertEqual(calcular_valor_dia(8, time(22, 30)), Decimal("125.00"))
        self.assertEqual(calcular_valor_dia(10, time(22, 30)), Decimal("125.00"))
        self.assertEqual(calcular_valor_dia(4, time(22, 15)), Decimal("62.50"))
        self.assertEqual(calcular_valor_dia(5, time(22, 15)), Decimal("62.50"))

    def test_fora_da_faixa_nao_gera_valor(self):
        self.assertIsNone(calcular_valor_dia(3, time(11, 0)))

    def test_exemplo_cinco_dias_440(self):
        dias = [
            _dia(8, time(17, 0), data=date(2026, 10, 1)),
            _dia(8, time(17, 0), data=date(2026, 10, 2)),
            _dia(8, time(17, 0), data=date(2026, 10, 3)),
            _dia(4, time(12, 0), data=date(2026, 10, 4)),
            _dia(4, time(12, 0), data=date(2026, 10, 5)),
        ]
        resumo = montar_resumo_pagamento(dias)
        self.assertEqual(resumo.total_dias, 5)
        self.assertEqual(resumo.dias_6_a_8, 3)
        self.assertEqual(resumo.dias_4, 2)
        self.assertEqual(resumo.valor_total, Decimal("440.00"))
        self.assertEqual(resumo.valor_total_formatado, "R$ 440,00")

    def test_exemplo_oito_dias_770(self):
        dias = [_dia(8, time(17, 0), data=date(2026, 10, d)) for d in range(1, 7)]
        dias += [
            _dia(4, time(12, 0), data=date(2026, 10, 7)),
            _dia(4, time(12, 0), data=date(2026, 10, 8)),
        ]
        resumo = montar_resumo_pagamento(dias)
        self.assertEqual(resumo.total_dias, 8)
        self.assertEqual(resumo.dias_6_a_8, 6)
        self.assertEqual(resumo.dias_4, 2)
        self.assertEqual(resumo.dias_apos_22, 0)
        self.assertEqual(resumo.valor_total, Decimal("770.00"))

    def test_um_dia_de_8h_apos_22_nao_usa_total_de_horas(self):
        dias = [
            _dia(8, time(22, 30), hora_inicio=time(14, 30), data=date(2026, 10, 1)),
            _dia(8, time(21, 0), hora_inicio=time(13, 0), data=date(2026, 10, 2)),
            _dia(4, time(20, 0), hora_inicio=time(16, 0), data=date(2026, 10, 3)),
        ]
        resumo = montar_resumo_pagamento(dias)
        self.assertEqual(resumo.dias_apos_22, 1)
        self.assertEqual(resumo.dias_ate_22, 2)
        self.assertEqual(resumo.valor_total, Decimal("110.00") + Decimal("125.00") + Decimal("55.00"))
        self.assertEqual(resumo.itens[0].valor_formatado, "R$ 125,00")
        self.assertEqual(resumo.itens[2].valor_formatado, "R$ 55,00")


class PrevisaoExcelTests(SimpleTestCase):
    def test_j26_continua_somando_o_valor_total(self):
        from core.services.excel_template import gerar_excel_previsao

        arquivo = gerar_excel_previsao(
            contexto={
                "numero_documento": "001/2026",
                "data_documento": "01/10/2026",
                "data_vencimento": "03/10/2026",
                "nome": "Ana Silva",
                "nome_curto": "Ana Silva",
                "empresa": "SESC",
                "valor_total": Decimal("440.00"),
                "valor_total_formatado": "R$ 440,00",
            },
            numero_documento="001/2026",
        )
        self.assertEqual(arquivo.name, "PP 001 - Ana Silva.xlsx")
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with zipfile.ZipFile(BytesIO(arquivo.read())) as planilha:
            self.assertIn("xl/drawings/drawing1.xml", planilha.namelist())
            root = etree.fromstring(planilha.read("xl/worksheets/sheet1.xml"))
        celulas = {c.get("r"): c for c in root.xpath("//m:c", namespaces=ns)}
        j17 = celulas["J17"]
        j26 = celulas["J26"]
        self.assertNotEqual(j17.get("t"), "s")
        self.assertEqual(j17.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v").text, "440.00")
        formula = j26.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f").text
        self.assertEqual(formula, "SUM(J16:J20)")
        self.assertEqual(
            j26.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}v").text,
            "440.00",
        )

    def test_rpa_grava_dias_e_valor_sem_apagar_formulas(self):
        from openpyxl import load_workbook

        from core.services.excel_template import gerar_excel_rpa

        arquivo = gerar_excel_rpa(
            contexto={
                "nome": "Ana Silva",
                "cpf": "000.000.000-00",
                "chave_pix": "ana@email.com",
                "numero_documento": "001/2026",
                "n_doc_sem_ano": "001",
                "dia_ano_doc": "01/2026",
                "data_documento": "01/10/2026",
                "quantidade_dias": 40,
                "valor_total": Decimal("440.00"),
            },
            numero_documento="001/2026",
        )
        conteudo = arquivo.read()
        with zipfile.ZipFile(BytesIO(conteudo)) as pacote:
            xml = pacote.read("xl/worksheets/sheet1.xml").decode()
        self.assertIn('<c r="N9"', xml)
        self.assertRegex(xml, r'<c r="N9"[^>]*>\s*<v>32</v>')
        planilha = load_workbook(BytesIO(conteudo))["Default"]
        self.assertEqual(planilha["N9"].number_format, "dd/mm/yyyy")
        self.assertEqual(planilha["N10"].value, 440)
        self.assertEqual(planilha["N10"].data_type, "n")
        self.assertEqual(planilha["N11"].value, "=N10*25%")
        self.assertEqual(planilha["N12"].value, "=N10-N11")
        self.assertEqual(planilha["N13"].value, "=(N10*N10)/N12")
        self.assertEqual(planilha["N14"].value, "=N13")
        self.assertEqual(planilha["F18"].value, "=N14")

    def test_rpa_usa_numero_proprio(self):
        from openpyxl import load_workbook

        from core.services.excel_template import gerar_excel_rpa

        arquivo = gerar_excel_rpa(
            contexto={
                "nome": "Ana Silva",
                "numero_documento": "013/2026",
                "n_doc_sem_ano": "013",
                "numero_rpa": "026/2026",
                "dia_ano_doc": "01/2026",
                "data_documento": "01/10/2026",
                "quantidade_dias": 3,
                "valor_total": Decimal("110.00"),
            },
            numero_documento="013/2026",
        )
        self.assertEqual(arquivo.name, "RPA 026 - Ana Silva.xlsx")
        conteudo = arquivo.read()
        with zipfile.ZipFile(BytesIO(conteudo)) as pacote:
            xml = pacote.read("xl/worksheets/sheet1.xml").decode()
        self.assertRegex(xml, r'<c r="N9"[^>]*>\s*<v>4</v>')
        planilha = load_workbook(BytesIO(conteudo))["Default"]
        self.assertEqual(planilha["N9"].number_format, "dd/mm/yyyy")
        self.assertEqual(planilha["N11"].value, "=N10*25%")
        textos = [
            str(cell.value)
            for row in planilha.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
        ]
        self.assertTrue(any("026" in texto for texto in textos))
        self.assertFalse(any(texto.strip() == "013" or "013/2026" in texto for texto in textos))


class AtestoAssinadoAnexoTests(TestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model

        from documentos.models import Documento
        from empresas.models import Empresa
        from pessoas.models import Pessoa
        from servicos.models import Servico

        self.user = get_user_model().objects.create_user("anexo", password="admin123")
        self.client.login(username="anexo", password="admin123")
        pessoa = Pessoa.objects.create(nome="Maria", cpf="52998224725", chave_pix="maria@email.com")
        empresa = Empresa.objects.create(nome="Empresa B")
        self.servico = Servico.objects.create(
            numero=91,
            ano=2026,
            numero_documento="091/2026",
            numero_rpa=91,
            ano_rpa=2026,
            numero_rpa_documento="R091/2026",
            pessoa=pessoa,
            empresa=empresa,
            data_servico=date(2026, 9, 1),
            motivo_contratacao="Cobertura",
            funcao="Assistente",
        )
        self.documento = Documento.objects.create(
            tipo_documento="DOCUMENTO_1",
            numero_documento=self.servico.numero_documento,
            pessoa=pessoa,
            servico=self.servico,
            conteudo_renderizado="CI",
            data_documento=date(2026, 9, 1),
        )

    def test_ficha_mostra_o_botao(self):
        resposta = self.client.get(reverse("servicos:detail", args=[self.servico.pk]))
        self.assertContains(resposta, "Anexar atesto assinado")
        documento = self.client.get(reverse("documentos:detail", args=[self.documento.pk]))
        self.assertContains(documento, "Anexar atesto assinado")

    def test_anexa_pdf_ou_imagem_no_servico(self):
        resposta = self.client.post(
            reverse("servicos:anexar_atesto", args=[self.servico.pk]),
            {"atesto_assinado": _png_minimo(), "next": reverse("servicos:detail", args=[self.servico.pk])},
        )
        self.assertRedirects(resposta, reverse("servicos:detail", args=[self.servico.pk]))
        self.servico.refresh_from_db()
        self.assertTrue(self.servico.atesto_assinado.name.endswith(".png"))
        ficha = self.client.get(reverse("servicos:detail", args=[self.servico.pk]))
        self.assertContains(ficha, "Atesto assinado anexado")

    def test_rejeita_arquivo_que_nao_e_pdf_nem_imagem(self):
        resposta = self.client.post(
            reverse("servicos:anexar_atesto", args=[self.servico.pk]),
            {"atesto_assinado": SimpleUploadedFile("nota.txt", b"texto", content_type="text/plain")},
        )
        self.assertEqual(resposta.status_code, 302)
        self.servico.refresh_from_db()
        self.assertFalse(self.servico.atesto_assinado)
        self.assertContains(
            self.client.get(reverse("servicos:detail", args=[self.servico.pk])),
            "Anexe um PDF ou uma imagem",
        )
