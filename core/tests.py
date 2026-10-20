from datetime import date, time
from decimal import Decimal
from concurrent.futures import ThreadPoolExecutor

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase, TransactionTestCase
from rest_framework.test import APIClient

from core.models import SequenciaAnual
from core.services.numeracao import gerar_proximo_numero
from core.services.servico_documentos import (
    ServicoDocumentoError,
    criar_servico_com_documentos,
    editar_servico_com_documentos,
)
from core.services.templates_texto import renderizar_texto, validar_cpf
from empresas.models import Empresa
from pessoas.models import Pessoa
from textos_padroes.models import TextoPadrao


def _cpf_valido():
    return "52998224725"


class CpfTests(TestCase):
    def test_cpf_valido(self):
        self.assertTrue(validar_cpf(_cpf_valido()))

    def test_cpf_invalido(self):
        self.assertFalse(validar_cpf("11111111111"))


class PessoaTests(TestCase):
    def test_cadastro_pessoa(self):
        p = Pessoa.objects.create(
            nome="João da Silva",
            cpf=_cpf_valido(),
            chave_pix="joao@email.com",
        )
        self.assertEqual(p.cpf, _cpf_valido())

    def test_cpf_invalido_bloqueado(self):
        with self.assertRaises(ValidationError):
            Pessoa.objects.create(
                nome="X",
                cpf="123",
                chave_pix="x@email.com",
            )


class EmpresaTextoTests(TestCase):
    def test_empresa_e_texto(self):
        e = Empresa.objects.create(nome="Empresa A")
        t = TextoPadrao.objects.create(
            nome="Previsão",
            conteudo="Olá {{ nome }} / [ CPF ]",
            tipo_documento="DOCUMENTO_1",
        )
        self.assertEqual(e.nome, "Empresa A")
        rendered = renderizar_texto(t.conteudo, {"nome": "Ana", "cpf": "1"})
        self.assertIn("Ana", rendered)
        self.assertIn("1", rendered)

    def test_formatacao_dias_horarios_iguais(self):
        from types import SimpleNamespace

        from core.services.templates_texto import formatar_dias_trabalhados

        dias = [
            SimpleNamespace(
                data=date(2026, 9, 5),
                hora_inicio=time(18, 0),
                hora_fim=time(22, 0),
            ),
            SimpleNamespace(
                data=date(2026, 9, 6),
                hora_inicio=time(18, 0),
                hora_fim=time(22, 0),
            ),
        ]
        texto = formatar_dias_trabalhados(dias)
        self.assertEqual(
            texto,
            "05 de Setembro até 06 de Setembro das 18:00h às 22:00h.",
        )

    def test_formatacao_dias_corridos_horarios_agrupados(self):
        from types import SimpleNamespace

        from core.services.templates_texto import formatar_dias_trabalhados

        dias = [
            SimpleNamespace(data=date(2026, 9, 5), hora_inicio=time(8, 0), hora_fim=time(17, 0)),
            SimpleNamespace(data=date(2026, 9, 6), hora_inicio=time(8, 0), hora_fim=time(10, 0)),
            SimpleNamespace(data=date(2026, 9, 7), hora_inicio=time(8, 0), hora_fim=time(17, 0)),
            SimpleNamespace(data=date(2026, 9, 8), hora_inicio=time(8, 0), hora_fim=time(10, 0)),
        ]
        texto = formatar_dias_trabalhados(dias)
        self.assertEqual(
            texto,
            "05 de Setembro até 08 de Setembro 05/09 e 07/09: 08:00h às 17:00h\n"
            "06/09 e 08/09: 08:00h às 10:00h.",
        )

    def test_formatacao_dias_quebrados_horarios_iguais(self):
        from types import SimpleNamespace

        from core.services.templates_texto import formatar_dias_trabalhados

        dias = [
            SimpleNamespace(data=date(2026, 9, 5), hora_inicio=time(8, 0), hora_fim=time(17, 0)),
            SimpleNamespace(data=date(2026, 9, 7), hora_inicio=time(8, 0), hora_fim=time(17, 0)),
        ]
        texto = formatar_dias_trabalhados(dias)
        self.assertEqual(texto, "05/09 e 07/09 das 08:00h às 17:00h.")

    def test_formatacao_dias_quebrados_horarios_divergentes(self):
        from types import SimpleNamespace

        from core.services.templates_texto import formatar_dias_trabalhados

        dias = [
            SimpleNamespace(data=date(2026, 9, 5), hora_inicio=time(8, 0), hora_fim=time(17, 0)),
            SimpleNamespace(data=date(2026, 9, 7), hora_inicio=time(9, 0), hora_fim=time(18, 0)),
        ]
        texto = formatar_dias_trabalhados(dias)
        self.assertEqual(
            texto,
            "05/09: 08:00h às 17:00h\n07/09: 09:00h às 18:00h.",
        )

    def test_formatacao_conjuntos_em_linhas_separadas(self):
        from types import SimpleNamespace

        from core.services.templates_texto import formatar_dias_trabalhados

        dias = [
            SimpleNamespace(data=date(2026, 9, 17), hora_inicio=time(11, 0), hora_fim=time(20, 0)),
            SimpleNamespace(data=date(2026, 9, 18), hora_inicio=time(6, 0), hora_fim=time(15, 0)),
            SimpleNamespace(data=date(2026, 9, 19), hora_inicio=time(8, 0), hora_fim=time(12, 0)),
            SimpleNamespace(data=date(2026, 9, 21), hora_inicio=time(6, 0), hora_fim=time(15, 0)),
            SimpleNamespace(data=date(2026, 9, 22), hora_inicio=time(6, 0), hora_fim=time(15, 0)),
        ]
        texto = formatar_dias_trabalhados(dias)
        self.assertEqual(
            texto,
            "17/09: 11:00h às 20:00h\n"
            "18/09, 21/09 e 22/09: 06:00h às 15:00h\n"
            "19/09: 08:00h às 12:00h.",
        )
        self.assertEqual(texto.count("\n"), 2)

    def test_ci_renderiza_quebra_de_linha_entre_horarios(self):
        from io import BytesIO

        from docx import Document

        from core.services.templates_texto import formatar_dias_trabalhados
        from core.services.word_template import gerar_documento_word
        from types import SimpleNamespace

        dias = [
            SimpleNamespace(data=date(2026, 9, 17), hora_inicio=time(11, 0), hora_fim=time(20, 0)),
            SimpleNamespace(data=date(2026, 9, 18), hora_inicio=time(6, 0), hora_fim=time(15, 0)),
            SimpleNamespace(data=date(2026, 9, 19), hora_inicio=time(8, 0), hora_fim=time(12, 0)),
            SimpleNamespace(data=date(2026, 9, 21), hora_inicio=time(6, 0), hora_fim=time(15, 0)),
            SimpleNamespace(data=date(2026, 9, 22), hora_inicio=time(6, 0), hora_fim=time(15, 0)),
        ]
        arquivo = gerar_documento_word(
            contexto={
                "nome": "Ana Silva",
                "cpf": "000.000.000-00",
                "chave_pix": "ana@email.com",
                "dias_trabalhados": formatar_dias_trabalhados(dias),
                "motivo_contratacao": "Cobertura",
                "funcao": "Porteiro",
                "numero_documento": "014/2026",
                "data_documento": "15/10/2026",
                "empresa": "SESC",
            },
            numero_documento="014/2026",
            versao=1,
        )
        documento = Document(BytesIO(arquivo.read()))
        paragrafo = next(p for p in documento.paragraphs if "17/09" in p.text)
        quebras = paragrafo._p.findall(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br"
        )
        self.assertGreaterEqual(len(quebras), 2)
        self.assertIn("17/09: 11:00h às 20:00h", paragrafo.text)
        self.assertIn("18/09, 21/09 e 22/09: 06:00h às 15:00h", paragrafo.text)
        self.assertIn("19/09: 08:00h às 12:00h.", paragrafo.text)

    def test_nome_curto_e_data_vencimento(self):
        from core.services.templates_texto import (
            adicionar_dias_uteis,
            extrair_nome_curto,
        )

        self.assertEqual(extrair_nome_curto("João da Silva"), "João Silva")
        self.assertEqual(extrair_nome_curto("Maria"), "Maria")
        # quinta 10/09/2026 + 1 útil = sexta 11/09/2026
        self.assertEqual(
            adicionar_dias_uteis(date(2026, 9, 10), 1),
            date(2026, 9, 11),
        )
        # sexta 11/09/2026 + 1 útil = segunda 14/09/2026
        self.assertEqual(
            adicionar_dias_uteis(date(2026, 9, 11), 1),
            date(2026, 9, 14),
        )
        # sábado e domingo também caem na segunda
        self.assertEqual(
            adicionar_dias_uteis(date(2026, 9, 12), 1),
            date(2026, 9, 14),
        )
        self.assertEqual(
            adicionar_dias_uteis(date(2026, 9, 13), 1),
            date(2026, 9, 14),
        )
        # quinta 10/09/2026 + 2 úteis = segunda 14/09/2026
        self.assertEqual(
            adicionar_dias_uteis(date(2026, 9, 10), 2),
            date(2026, 9, 14),
        )
        # sexta 11/09/2026 + 2 úteis = terça 15/09/2026
        self.assertEqual(
            adicionar_dias_uteis(date(2026, 9, 11), 2),
            date(2026, 9, 15),
        )
        # sábado e domingo + 2 úteis também caem na terça
        self.assertEqual(
            adicionar_dias_uteis(date(2026, 9, 12), 2),
            date(2026, 9, 15),
        )
        self.assertEqual(
            adicionar_dias_uteis(date(2026, 9, 13), 2),
            date(2026, 9, 15),
        )

    def test_n_doc_sem_ano_e_dia_ano_doc(self):
        from core.services.templates_texto import (
            extrair_n_doc_sem_ano,
            formatar_dia_ano_doc,
        )

        self.assertEqual(extrair_n_doc_sem_ano("001/2026"), "001")
        self.assertEqual(extrair_n_doc_sem_ano("12/2026"), "12")
        self.assertEqual(formatar_dia_ano_doc(date(2026, 9, 15)), "15/2026")
        self.assertEqual(formatar_dia_ano_doc("05/09/2026"), "05/2026")

    def test_atesto_preenche_variaveis_e_nome_do_arquivo(self):
        from docx import Document
        from io import BytesIO

        from core.services.word_template import gerar_documento_atesto

        arquivo = gerar_documento_atesto(
            contexto={
                "nome": "Ana Silva",
                "cpf": "000.000.000-00",
                "empresa": "SESC Arcoverde",
                "funcao": "Porteiro",
                "dias_trabalhados": "05 de Setembro até 06 de Setembro das 08:00h às 17:00h.",
                "horarios_trabalhados": "08:00h às 17:00h",
                "observacoes": "Cobertura",
                "data_documento": "15/09/2026",
                "numero_documento": "013/2026",
            },
            numero_documento="013/2026",
        )
        self.assertEqual(arquivo.name, "ATESTO 013 - Ana Silva.docx")
        documento = Document(BytesIO(arquivo.read()))
        texto = "\n".join(p.text for p in documento.paragraphs)
        self.assertIn("UNIDADE: SESC Arcoverde", texto)
        self.assertIn("Nome do Prestador de Serviço: Ana Silva", texto)
        self.assertIn("CPF: 000.000.000-00", texto)
        self.assertIn("Função/Atividade: Porteiro", texto)
        self.assertIn("Horário: 08:00h às 17:00h", texto)
        self.assertIn("Observações (se houver): Cobertura", texto)
        self.assertIn(
            "Data da Prestação do Serviço: 05 de Setembro até 06 de Setembro das 08:00h às 17:00h.",
            texto,
        )
        self.assertIn("Data: 15/09/2026", texto)
        self.assertNotIn("{{", texto)
        unidade = next(p for p in documento.paragraphs if p.text.startswith("UNIDADE:"))
        self.assertTrue(unidade.runs[0].bold)
        self.assertEqual(unidade.runs[1].text, "SESC Arcoverde")
        self.assertIs(unidade.runs[1].bold, False)
        horario = next(p for p in documento.paragraphs if p.text.startswith("Horário:"))
        self.assertIs(horario.runs[1].bold, False)

    def test_datas_trabalhadas_sem_horario(self):
        from types import SimpleNamespace

        from core.services.templates_texto import formatar_datas_trabalhadas

        corridos = [
            SimpleNamespace(data=date(2026, 9, 5)),
            SimpleNamespace(data=date(2026, 9, 6)),
        ]
        self.assertEqual(
            formatar_datas_trabalhadas(corridos),
            "05 de Setembro até 06 de Setembro",
        )
        quebrados = [
            SimpleNamespace(data=date(2026, 8, 25)),
            SimpleNamespace(data=date(2026, 8, 26)),
            SimpleNamespace(data=date(2026, 8, 30)),
        ]
        self.assertEqual(
            formatar_datas_trabalhadas(quebrados),
            "25/08, 26/08 e 30/08",
        )


class NumeracaoTests(TransactionTestCase):
    def test_sequencial(self):
        n1, a1, f1 = gerar_proximo_numero(SequenciaAnual.Tipo.SERVICO, ano=2026)
        n2, a2, f2 = gerar_proximo_numero(SequenciaAnual.Tipo.SERVICO, ano=2026)
        self.assertEqual(f1, "001/2026")
        self.assertEqual(f2, "002/2026")
        self.assertEqual(n2, n1 + 1)

    def test_concorrencia(self):
        from django.db import connection

        if connection.vendor != "postgresql":
            self.skipTest("Teste de concorrência requer PostgreSQL (select_for_update).")

        SequenciaAnual.objects.create(
            tipo=SequenciaAnual.Tipo.SERVICO, ano=2030, ultimo_numero=0
        )

        def worker(_):
            return gerar_proximo_numero(SequenciaAnual.Tipo.SERVICO, ano=2030)[0]

        with ThreadPoolExecutor(max_workers=5) as pool:
            numeros = list(pool.map(worker, range(10)))
        self.assertEqual(len(numeros), len(set(numeros)))
        self.assertEqual(sorted(numeros), list(range(1, 11)))


class ServicoFlowTests(TestCase):
    def setUp(self):
        self.pessoa = Pessoa.objects.create(
            nome="Maria",
            cpf=_cpf_valido(),
            chave_pix="maria@email.com",
        )
        self.empresa = Empresa.objects.create(nome="Empresa B")
        self.textos = {}
        for i, tipo in enumerate(("DOCUMENTO_1", "DOCUMENTO_2", "DOCUMENTO_3"), start=1):
            self.textos[tipo] = TextoPadrao.objects.create(
                nome=f"Texto {i}",
                conteudo=f"Doc {i} {{ nome }} {{ numero_documento }} {{ total_horas }}",
                tipo_documento=tipo,
            )
        self.user = get_user_model().objects.create_user("admin", password="admin123")

    def test_criar_servico_com_3_documentos_e_pdf(self):
        servico = criar_servico_com_documentos(
            pessoa=self.pessoa,
            empresa=self.empresa,
            motivo_contratacao="Cobertura",
            funcao="Assistente",
            data_servico=date(2026, 9, 1),
            dias_trabalhados=[
                {"data": date(2026, 9, 1), "hora_inicio": "08:00", "hora_fim": "16:00"},
                {"data": date(2026, 9, 2), "hora_inicio": "08:00", "hora_fim": "14:00"},
            ],
            textos_por_tipo=self.textos,
            usuario=self.user,
        )
        self.assertEqual(servico.documentos.count(), 3)
        self.assertTrue(all(d.arquivo_pdf for d in servico.documentos.all()))
        self.assertTrue(
            all(str(d.arquivo_pdf.name).endswith(".docx") for d in servico.documentos.all())
        )
        self.assertEqual(servico.dias_trabalhados.count(), 2)
        self.assertEqual(float(servico.total_horas), 14.0)

    def test_ci_e_rpa_tem_sequencias_independentes(self):
        from core.models import SequenciaAnual
        from core.services.numeracao import definir_ultimo_numero

        definir_ultimo_numero(SequenciaAnual.Tipo.SERVICO, 12, ano=2026)
        definir_ultimo_numero(SequenciaAnual.Tipo.RPA, 25, ano=2026)
        primeiro = criar_servico_com_documentos(
            pessoa=self.pessoa,
            empresa=self.empresa,
            motivo_contratacao="Cobertura",
            funcao="Assistente",
            data_servico=date(2026, 9, 1),
            dias_trabalhados=[
                {"data": date(2026, 9, 1), "hora_inicio": "08:00", "hora_fim": "16:00"}
            ],
            textos_por_tipo={"DOCUMENTO_1": self.textos["DOCUMENTO_1"]},
            usuario=self.user,
        )
        segundo = criar_servico_com_documentos(
            pessoa=self.pessoa,
            empresa=self.empresa,
            motivo_contratacao="Cobertura",
            funcao="Assistente",
            data_servico=date(2026, 9, 2),
            dias_trabalhados=[
                {"data": date(2026, 9, 2), "hora_inicio": "08:00", "hora_fim": "16:00"}
            ],
            textos_por_tipo={"DOCUMENTO_1": self.textos["DOCUMENTO_1"]},
            usuario=self.user,
        )
        self.assertEqual(primeiro.numero_documento, "013/2026")
        self.assertEqual(primeiro.numero_rpa_documento, "026/2026")
        self.assertEqual(segundo.numero_documento, "014/2026")
        self.assertEqual(segundo.numero_rpa_documento, "027/2026")

    def test_edicao_gera_historico(self):
        servico = criar_servico_com_documentos(
            pessoa=self.pessoa,
            empresa=self.empresa,
            motivo_contratacao="Cobertura",
            funcao="Assistente",
            data_servico=date(2026, 9, 1),
            dias_trabalhados=[{"data": date(2026, 9, 1), "hora_inicio": "08:00", "hora_fim": "16:00"}],
            textos_por_tipo=self.textos,
            usuario=self.user,
        )
        editar_servico_com_documentos(
            servico=servico,
            motivo_contratacao="Novo motivo",
            funcao="Analista",
            data_servico=date(2026, 9, 1),
            dias_trabalhados=[{"data": date(2026, 9, 3), "hora_inicio": "09:00", "hora_fim": "13:00"}],
            textos_por_tipo=self.textos,
            usuario=self.user,
            motivo_alteracao="Ajuste de horas",
        )
        doc = servico.documentos.first()
        self.assertEqual(doc.versao_atual, 2)
        self.assertEqual(doc.versoes.count(), 2)

    def test_criar_servico_somente_documento_1(self):
        servico = criar_servico_com_documentos(
            pessoa=self.pessoa,
            empresa=self.empresa,
            motivo_contratacao="Cobertura",
            funcao="Assistente",
            data_servico=date(2026, 9, 1),
            dias_trabalhados=[
                {"data": date(2026, 9, 1), "hora_inicio": "08:00", "hora_fim": "16:00"}
            ],
            textos_por_tipo={"DOCUMENTO_1": self.textos["DOCUMENTO_1"]},
            usuario=self.user,
        )
        self.assertEqual(servico.documentos.count(), 1)
        self.assertEqual(
            servico.documentos.get().tipo_documento, "DOCUMENTO_1"
        )

    def test_assinatura_define_vencimento_em_dia_util(self):
        from unittest.mock import patch

        from core.services.servico_documentos import montar_contexto_de_servico
        from servicos.models import Servico

        textos = {
            "DOCUMENTO_1": TextoPadrao.objects.create(
                nome="Vencimento",
                conteudo="Vence em {{ data_vencimento }}",
                tipo_documento="DOCUMENTO_1",
            )
        }
        dias = [
            {"data": date(2026, 9, 1), "hora_inicio": "08:00", "hora_fim": "16:00"}
        ]
        sexta = date(2026, 9, 11)
        with patch(
            "core.services.servico_documentos.timezone.localdate",
            return_value=sexta,
        ):
            walquiria = criar_servico_com_documentos(
                pessoa=self.pessoa,
                empresa=self.empresa,
                motivo_contratacao="Cobertura",
                funcao="Assistente",
                data_servico=date(2026, 9, 1),
                assinatura=Servico.Assinatura.WALQUIRIA,
                dias_trabalhados=dias,
                textos_por_tipo=textos,
                usuario=self.user,
            )
            asaphe = criar_servico_com_documentos(
                pessoa=self.pessoa,
                empresa=self.empresa,
                motivo_contratacao="Cobertura",
                funcao="Assistente",
                data_servico=date(2026, 9, 1),
                assinatura=Servico.Assinatura.ASAPHE,
                dias_trabalhados=dias,
                textos_por_tipo=textos,
                usuario=self.user,
            )

        self.assertEqual(walquiria.assinatura, Servico.Assinatura.WALQUIRIA)
        self.assertEqual(asaphe.assinatura, Servico.Assinatura.ASAPHE)
        self.assertEqual(walquiria.dias_uteis_vencimento(), 1)
        self.assertEqual(asaphe.dias_uteis_vencimento(), 2)
        self.assertIn("14/09/2026", walquiria.documentos.get().conteudo_renderizado)
        self.assertIn("15/09/2026", asaphe.documentos.get().conteudo_renderizado)
        self.assertEqual(
            montar_contexto_de_servico(walquiria)["data_vencimento"],
            "14/09/2026",
        )
        self.assertEqual(
            montar_contexto_de_servico(asaphe)["data_vencimento"],
            "15/09/2026",
        )

        quinta = date(2026, 9, 10)
        with patch(
            "core.services.servico_documentos.timezone.localdate",
            return_value=quinta,
        ):
            editar_servico_com_documentos(
                servico=asaphe,
                motivo_contratacao="Cobertura",
                funcao="Assistente",
                data_servico=date(2026, 9, 1),
                assinatura=Servico.Assinatura.ASAPHE,
                dias_trabalhados=dias,
                textos_por_tipo=textos,
                usuario=self.user,
            )
        asaphe.refresh_from_db()
        self.assertEqual(
            montar_contexto_de_servico(asaphe)["data_vencimento"],
            "14/09/2026",
        )

    def test_horas_invalidas(self):
        with self.assertRaises(ServicoDocumentoError):
            criar_servico_com_documentos(
                pessoa=self.pessoa,
                empresa=self.empresa,
                motivo_contratacao="X",
                funcao="Y",
                data_servico=date(2026, 9, 1),
                dias_trabalhados=[
                    {"data": date(2026, 9, 1), "hora_inicio": "08:00", "hora_fim": "08:00"}
                ],
                textos_por_tipo={"DOCUMENTO_1": self.textos["DOCUMENTO_1"]},
            )

    def test_horario_atravessa_meia_noite_sem_mudar_a_data(self):
        servico = criar_servico_com_documentos(
            pessoa=self.pessoa,
            empresa=self.empresa,
            motivo_contratacao="Plantão",
            funcao="Porteiro",
            data_servico=date(2026, 9, 17),
            dias_trabalhados=[
                {"data": date(2026, 9, 17), "hora_inicio": "18:00", "hora_fim": "06:00"}
            ],
            textos_por_tipo={"DOCUMENTO_1": self.textos["DOCUMENTO_1"]},
            usuario=self.user,
        )
        dia = servico.dias_trabalhados.get()
        self.assertEqual(dia.data, date(2026, 9, 17))
        self.assertEqual(servico.data_servico, date(2026, 9, 17))
        self.assertEqual(dia.quantidade_horas, Decimal("12.00"))
        self.assertEqual(servico.total_horas, Decimal("12.00"))

    def test_edicao_sugere_data_atual_e_busca_ignora_acento(self):
        from django.utils import timezone

        servico = criar_servico_com_documentos(
            pessoa=self.pessoa,
            empresa=self.empresa,
            motivo_contratacao="Cobertura",
            funcao="Assistente",
            data_servico=date(2026, 9, 1),
            dias_trabalhados=[
                {"data": date(2026, 9, 1), "hora_inicio": "08:00", "hora_fim": "16:00"}
            ],
            textos_por_tipo={"DOCUMENTO_1": self.textos["DOCUMENTO_1"]},
            usuario=self.user,
        )
        Empresa.objects.create(nome="José Alimentos")
        self.client.login(username="admin", password="admin123")
        edicao = self.client.get(f"/servicos/{servico.pk}/editar/")
        html = edicao.content.decode()
        hoje = timezone.localdate().isoformat()
        self.assertIn(f'value="{hoje}"', html)
        self.assertNotIn('value="2026-09-01"', html)

        busca = self.client.get("/servicos/busca-cadastro/", {"tipo": "empresa", "q": "jose"})
        self.assertEqual(busca.status_code, 200)
        nomes = [item["nome"] for item in busca.json()["results"]]
        self.assertIn("José Alimentos", nomes)
        self.assertNotIn("Empresa B", nomes)

        pessoa = self.client.get("/servicos/busca-cadastro/", {"tipo": "pessoa", "q": "MARIA"})
        self.assertEqual([item["id"] for item in pessoa.json()["results"]], [self.pessoa.pk])


class AssinaturaFormTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("formuser", password="form123")
        self.client.login(username="formuser", password="form123")

    def test_criacao_exibe_select_de_assinatura(self):
        response = self.client.get("/servicos/novo/")
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn("Assinatura", html)
        self.assertIn('name="assinatura"', html)
        self.assertIn('value="ASAPHE"', html)
        self.assertIn(">Asaphe<", html)
        self.assertIn('value="WALQUIRIA"', html)
        self.assertIn(">Walquiria<", html)
        self.assertIn("WALQUIRIA\" selected", html)
        from django.utils import timezone

        self.assertIn(f'value="{timezone.localdate().isoformat()}"', html)
        self.assertIn('data-busca-cadastro="empresa"', html)
        self.assertIn('data-busca-cadastro="pessoa"', html)
        self.assertIn("data-horario", html)
        self.assertNotIn('type="time"', html)


class ApiPermissaoTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = get_user_model().objects.create_user("api", password="api123")

    def test_api_exige_auth(self):
        resp = self.client.get("/api/pessoas/")
        self.assertIn(resp.status_code, (401, 403))

    def test_api_pessoa_autenticada(self):
        self.client.force_authenticate(self.user)
        resp = self.client.post(
            "/api/pessoas/",
            {
                "nome": "API User",
                "cpf": _cpf_valido(),
                "chave_pix": "pix@email.com",
                "ativo": True,
            },
            format="json",
        )
        self.assertEqual(resp.status_code, 201)
        list_resp = self.client.get("/api/pessoas/")
        self.assertEqual(list_resp.status_code, 200)
        self.assertNotIn("chave_pix", list_resp.json()["results"][0])
