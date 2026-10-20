import re
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.core.files.base import ContentFile
from docx import Document

from core.services.nomes_arquivo import nome_arquivo_download


TEMPLATE_DOCX = (
    Path(settings.BASE_DIR) / "media" / "documentos_manuais" / "ci_padrao.docx"
)
TEMPLATE_ATESTO = (
    Path(settings.BASE_DIR) / "media" / "documentos_manuais" / "atesto_template.docx"
)
PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")


def _substituir_texto(texto: str, contexto: dict) -> str:
    def _sub(match: re.Match) -> str:
        chave = match.group(1).strip().lower()
        valor = contexto.get(chave, "")
        return "" if valor is None else str(valor)

    return PLACEHOLDER_RE.sub(_sub, texto or "")


def _substituir_em_paragrafo(paragraph, contexto: dict) -> None:
    original = paragraph.text
    if "{{" not in original:
        return
    novo = _substituir_texto(original, contexto)
    if novo == original:
        return
    if paragraph.runs:
        paragraph.runs[0].text = novo
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(novo)


def _partir_variaveis(texto: str, contexto: dict) -> list[tuple[str, bool]]:
    partes: list[tuple[str, bool]] = []
    ultima = 0
    for match in PLACEHOLDER_RE.finditer(texto):
        if match.start() > ultima:
            partes.append((texto[ultima:match.start()], False))
        valor = contexto.get(match.group(1).strip().lower(), "")
        partes.append(("" if valor is None else str(valor), True))
        ultima = match.end()
    if ultima < len(texto):
        partes.append((texto[ultima:], False))
    return partes


def _run_depois(paragraph, run, texto: str, negrito: bool):
    novo = paragraph.add_run(texto)
    novo.bold = negrito
    novo.italic = run.italic
    novo.font.name = run.font.name
    novo.font.size = run.font.size
    run._element.addnext(novo._element)
    return novo


def _substituir_variavel_sem_negrito(paragraph, contexto: dict) -> None:
    """Troca só o trecho da variável e grava esse trecho sem negrito."""
    for run in list(paragraph.runs):
        if "{{" not in (run.text or ""):
            continue
        partes = _partir_variaveis(run.text, contexto)
        if not partes:
            continue
        primeiro, primeiro_e_variavel = partes[0]
        run.text = primeiro
        if primeiro_e_variavel:
            run.bold = False
        cursor = run
        for trecho, eh_variavel in partes[1:]:
            cursor = _run_depois(
                paragraph,
                cursor,
                trecho,
                False if eh_variavel else bool(run.bold),
            )


def _substituir_em_documento(document: Document, contexto: dict, *, sem_negrito: bool = False) -> None:
    paragrafos = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragrafos.extend(cell.paragraphs)
    for paragraph in paragrafos:
        if sem_negrito:
            _substituir_variavel_sem_negrito(paragraph, contexto)
        else:
            _substituir_em_paragrafo(paragraph, contexto)


def _gerar_docx(
    template: Path,
    contexto: dict,
    prefixo: str,
    numero_documento: str,
    *,
    sem_negrito: bool = False,
) -> ContentFile:
    if not template.exists():
        raise FileNotFoundError(f"Modelo Word não encontrado em {template}.")

    document = Document(str(template))
    _substituir_em_documento(document, contexto, sem_negrito=sem_negrito)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    nome = nome_arquivo_download(
        prefixo, numero_documento, contexto.get("nome", ""), "docx"
    )
    return ContentFile(buffer.read(), name=nome)


def gerar_documento_word(*, contexto: dict, numero_documento: str, versao: int) -> ContentFile:
    """
    Preenche o modelo CI (ci_padrao.docx) com as variáveis do contexto.

    Variáveis esperadas no modelo:
    nome, cpf, chave_pix, dias_trabalhados, motivo_contratacao, funcao,
    numero_documento, data_documento, empresa
    """
    return _gerar_docx(TEMPLATE_DOCX, contexto, "CI", numero_documento)


def gerar_documento_atesto(*, contexto: dict, numero_documento: str) -> ContentFile:
    """
    Preenche o atesto (atesto_template.docx).

    Variáveis: empresa, dias_trabalhados, nome, cpf, funcao,
    horarios_trabalhados, observacoes, data_documento.
    """
    return _gerar_docx(
        TEMPLATE_ATESTO, contexto, "ATESTO", numero_documento, sem_negrito=True
    )
