import re
import zipfile
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.core.files.base import ContentFile
from lxml import etree
from openpyxl import load_workbook

from core.services.nomes_arquivo import nome_arquivo_download
from core.services.templates_texto import extrair_n_doc_sem_ano

_SSML = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

MEDIA_MANUAIS = Path(settings.BASE_DIR) / "media" / "documentos_manuais"
TEMPLATE_PREVISAO_XLSX = MEDIA_MANUAIS / "modelo_previsaodepagamento.xlsx"
TEMPLATE_RPA_XLSX = MEDIA_MANUAIS / "modelo_RPA.xlsx"

PLACEHOLDER_RE = re.compile(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}")

# Partes XML onde as variáveis do modelo de previsão podem aparecer.
_XML_PARTS_PREVISAO = (
    "xl/sharedStrings.xml",
    "xl/worksheets/sheet1.xml",
)


def _valor_contexto(chave: str, contexto: dict) -> str:
    valor = contexto.get(chave, "")
    if valor is None:
        return ""
    return str(valor)


def _substituir_texto(texto: str, contexto: dict, *, escape_xml: bool = False) -> str:
    def _sub(match: re.Match) -> str:
        chave = match.group(1).strip().lower()
        valor = _valor_contexto(chave, contexto)
        if escape_xml:
            return (
                valor.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )
        return valor

    return PLACEHOLDER_RE.sub(_sub, texto or "")


def _numero_excel(contexto: dict) -> str:
    """Número aceito pelo Excel (ponto decimal), para entrar em SUM."""
    bruto = contexto.get("valor_total")
    if bruto not in (None, ""):
        return format(Decimal(str(bruto)), "f")
    texto = str(contexto.get("valor_total_formatado") or "0")
    limpo = (
        texto.replace("R$", "")
        .replace("\u00a0", "")
        .replace(" ", "")
        .replace(".", "")
        .replace(",", ".")
    )
    try:
        return format(Decimal(limpo or "0"), "f")
    except InvalidOperation:
        return "0"


def _indices_placeholder_valor(shared_xml: bytes) -> set[str]:
    root = etree.fromstring(shared_xml)
    indices = set()
    for indice, item in enumerate(root.findall(f"{{{_SSML}}}si")):
        texto = "".join(item.itertext()).strip()
        if texto == "{{ valor_total_formatado }}":
            indices.add(str(indice))
    return indices


def _celula_valor_para_numero(cell, numero: str) -> None:
    if "t" in cell.attrib:
        del cell.attrib["t"]
    for filho in list(cell):
        if etree.QName(filho).localname in ("v", "is"):
            cell.remove(filho)
    valor = etree.Element(f"{{{_SSML}}}v")
    valor.text = numero
    cell.append(valor)


def _atualizar_cache_somas(root, numero: str) -> None:
    for cell in root.findall(f".//{{{_SSML}}}c"):
        formula = cell.find(f"{{{_SSML}}}f")
        if formula is None or not formula.text:
            continue
        texto = formula.text.replace(" ", "")
        if texto not in ("SUM(J16:J20)", 'IF(J26="","",J26)'):
            continue
        cache = cell.find(f"{{{_SSML}}}v")
        if cache is None:
            cache = etree.SubElement(cell, f"{{{_SSML}}}v")
        cache.text = numero
        cell.set("t", "n")


def _aplicar_valor_numerico_na_soma(sheet_xml: bytes, shared_xml: bytes, numero: str) -> bytes:
    """
    Troca {{ valor_total_formatado }} por número na célula.
    SUM(J16:J20) ignora texto; o valor precisa ser numérico para J26 somar.
    """
    indices = _indices_placeholder_valor(shared_xml)
    if not indices:
        return sheet_xml
    root = etree.fromstring(sheet_xml)
    for cell in root.findall(f".//{{{_SSML}}}c"):
        if cell.get("t") != "s":
            continue
        valor = cell.find(f"{{{_SSML}}}v")
        if valor is None or valor.text not in indices:
            continue
        _celula_valor_para_numero(cell, numero)
    _atualizar_cache_somas(root, numero)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _preencher_previsao(template: Path, contexto: dict) -> bytes:
    numero = _numero_excel(contexto)
    raw = template.read_bytes()
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(raw), "r") as zin:
        shared = zin.read("xl/sharedStrings.xml")
        planilha = _aplicar_valor_numerico_na_soma(
            zin.read("xl/worksheets/sheet1.xml"), shared, numero
        )
        with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                if info.filename == "xl/worksheets/sheet1.xml":
                    data = planilha
                else:
                    data = zin.read(info.filename)
                if info.filename in _XML_PARTS_PREVISAO:
                    texto = data.decode("utf-8")
                    if "{{" in texto:
                        texto = _substituir_texto(texto, contexto, escape_xml=True)
                        data = texto.encode("utf-8")
                zout.writestr(info, data)
    out.seek(0)
    return out.read()


def gerar_excel_previsao(*, contexto: dict, numero_documento: str) -> ContentFile:
    """
    Preenche o modelo de previsão de pagamento.

    Usa edição direta do ZIP/XML para não destruir caixas de texto
    flutuantes (openpyxl remove shapes ao salvar).

    Variáveis no Excel:
    numero_documento, data_documento, data_vencimento, nome_curto, empresa,
    valor_total_formatado
    """
    if not TEMPLATE_PREVISAO_XLSX.exists():
        raise FileNotFoundError(
            f"Modelo Excel não encontrado em {TEMPLATE_PREVISAO_XLSX}. "
            "Converta modelo_previsaodepagamento.xls para .xlsx."
        )

    conteudo = _preencher_previsao(TEMPLATE_PREVISAO_XLSX, contexto)
    nome = nome_arquivo_download(
        "PP", numero_documento, contexto.get("nome", ""), "xlsx"
    )
    return ContentFile(conteudo, name=nome)


def _quantidade_dias_rpa(contexto: dict) -> int:
    """Dias trabalhados para N9, no máximo 31, para caber no formato de data."""
    try:
        dias = int(contexto.get("quantidade_dias") or 0)
    except (TypeError, ValueError):
        dias = 0
    return max(0, min(dias, 31))


def _preencher_quantidades_rpa(worksheet, contexto: dict) -> None:
    """
    N9 recebe os dias trabalhados como data em janeiro de 1900.

    3 dias viram 03/01/1900. Acima de 31, fica 31/01/1900.
    O mês e o ano permanecem 01/1900. N10 e as fórmulas seguintes não mudam.
    """
    n9 = worksheet["N9"]
    dias = _quantidade_dias_rpa(contexto)
    # Janeiro de 1900: o serial n aparece como o dia n-1 no LibreOffice,
    # por causa do dia 29/02/1900 que o Excel inventou. Somar 1 faz
    # 3 dias aparecerem como 03/01/1900.
    n9.value = dias + 1 if dias else 0
    n9.number_format = "dd/mm/yyyy"

    valor = contexto.get("valor_total")
    if valor in (None, ""):
        valor = 0
    n10 = worksheet["N10"]
    formato_valor = n10.number_format
    n10.value = Decimal(str(valor))
    n10.number_format = formato_valor


def _contexto_do_rpa(contexto: dict, numero_documento: str) -> tuple[dict, str]:
    """O RPA usa a sequência própria; CI e previsão continuam em numero_documento."""
    numero = contexto.get("numero_rpa") or numero_documento
    ajustado = dict(contexto)
    ajustado["numero_documento"] = numero
    ajustado["n_doc_sem_ano"] = extrair_n_doc_sem_ano(str(numero or ""))
    return ajustado, str(numero or "")


def gerar_excel_rpa(*, contexto: dict, numero_documento: str) -> ContentFile:
    """
    Preenche o modelo RPA.

    Variáveis principais:
    nome, cpf, chave_pix, data_documento, n_doc_sem_ano, dia_ano_doc

    O número impresso é numero_rpa, quando o contexto traz a sequência do RPA.
    """
    if not TEMPLATE_RPA_XLSX.exists():
        raise FileNotFoundError(
            f"Modelo RPA não encontrado em {TEMPLATE_RPA_XLSX}."
        )

    contexto, numero_rpa = _contexto_do_rpa(contexto, numero_documento)
    workbook = load_workbook(str(TEMPLATE_RPA_XLSX))
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value:
                    cell.value = _substituir_texto(cell.value, contexto)
        if worksheet.title != "IR-2024":
            _preencher_quantidades_rpa(worksheet, contexto)

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    nome = nome_arquivo_download(
        "RPA", numero_rpa, contexto.get("nome", ""), "xlsx"
    )
    return ContentFile(buffer.read(), name=nome)
